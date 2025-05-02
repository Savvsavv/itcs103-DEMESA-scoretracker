import tkinter as tk
from tkinter import messagebox
from openpyxl import load_workbook, Workbook

window = tk.Tk()
window.title("User Score Entry")
window.geometry('600x300')
window.configure(bg="gray25")

def create_excel():
    wb = Workbook()
    ws = wb.active
    ws.title = "score"
    ws.append(["Name", "Score", "Remarks"])
    wb.save("data.xlsx")

def edit_to_excel():
    if not crud():
        return
    nm = name_entry.get().strip()
    sc = int(score_entry.get())
    
    try:
        wb = load_workbook("data.xlsx")
        if "score" in wb.sheetnames:
            ws = wb["score"]
        else:
            ws = wb.active
    except FileNotFoundError:
        wb = Workbook()
        ws = wb.active
        ws.title = "score"
        ws.append(["Name", "Score", "Remarks"])

    if 75 <= sc <= 100:
        ws.append([nm, sc, "Passed"])
        messagebox.showinfo(title="passed", message="you passed the exam!")
    elif 0 <= sc <= 74:
        ws.append([nm, sc, "Failed"])
        messagebox.showinfo(title="failed", message="you failed the exam!")
    else:
        messagebox.showerror(title="Error Input", message="error, please enter a valid score between 0 to 100.")

    wb.save("data.xlsx")
    messagebox.showinfo(title="Success", message="your data has been saved successfully!")

def update_to_excel():
    if not crud():
        return
    nm = name_entry.get().strip()
    sc = int(score_entry.get())

    try:
        wb = load_workbook("data.xlsx")
        ws = wb["score"]

        for row in ws.iter_rows(min_row=2):
            row_val = row[0].value
            if row_val == nm:
                row[1].value = sc
                if 75 <= sc <= 100:
                    row[2].value = "Passed"
                elif 0 <= sc <= 74:
                    row[2].value = "Failed"
                wb.save("data.xlsx")
                messagebox.showinfo(title="Success", message="Updated Successfully")
                break
        else:
            messagebox.showerror(title="Error", message="not found!")
    except FileNotFoundError:
        messagebox.showerror(title="Error", message="File not found!")

def crud():
    try:
        score = int(score_entry.get())
        if 0 <= score <= 100:
            if score > 74:
                messagebox.showinfo(title="Passed", message="You passed the exam!")
            else:
                messagebox.showinfo(title="Failed", message="You failed the exam!")
            return True
        else:
            messagebox.showerror(title="Error", message="Input error, please enter a valid score between 0 and 100.")
            return False
    except ValueError:
        messagebox.showerror(title="Error", message="Input error, please enter a valid number!")
        return False

def show_data():
    wb = load_workbook("data.xlsx")
    ws = wb["score"]

    data_window = tk.Toplevel(window)
    data_window.title("score")
    


    for i, row in enumerate(ws.iter_rows(values_only=True)):
        for j, value in enumerate(row):
            label = tk.Label(data_window, text=value, borderwidth=1, relief="solid", padx=6, pady=3)
            label.grid(row=i, column=j)

# Main frame
mainframe = tk.Frame(window, bg="gray25", padx=20, pady=20)
mainframe.place(anchor="center", relx=0.5, rely=0.5)

# Title Label
title_label = tk.Label(mainframe, text="Student Score Entry", bg="gray25", fg="white",
                       font=("Arial", 14, "bold"))
title_label.grid(row=0, column=0, columnspan=3, pady=(0, 20))

# Labels
name_label = tk.Label(mainframe, text="Name:", bg="gray25", fg="white", font=("Arial", 12))
name_label.grid(row=1, column=0, sticky="e", pady=5)

score_label = tk.Label(mainframe, text="Score:", bg="gray25", fg="white", font=("Arial", 12))
score_label.grid(row=2, column=0, sticky="e", pady=5)

# Entry fields
name_entry = tk.Entry(mainframe, width=25)
name_entry.grid(row=1, column=1, columnspan=2, pady=5)

score_entry = tk.Entry(mainframe, width=25)
score_entry.grid(row=2, column=1, columnspan=2, pady=5)

# Buttons
enter_button = tk.Button(mainframe, text="Enter", command=edit_to_excel, width=10)
enter_button.grid(row=3, column=1, pady=10, padx=5, sticky="e")

update_button = tk.Button(mainframe, text="Update", command=update_to_excel, width=10)
update_button.grid(row=3, column=2, pady=10, padx=5, sticky="w")

showdata_button = tk.Button(mainframe, text="Show Data", command=show_data, width=22)
showdata_button.grid(row=4, column=1, columnspan=2, pady=10)

window.mainloop()
